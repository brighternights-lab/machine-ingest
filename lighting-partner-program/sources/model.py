from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb=Workbook()
DARK=PatternFill('solid',fgColor='0E111B'); HEAD=PatternFill('solid',fgColor='1A1F2E'); INP=PatternFill('solid',fgColor='1F2A44')
WH=Font(name='Arial',color='FFFFFF',bold=True,size=11); INK=Font(name='Arial',color='D8DBE3',size=10); PINK=Font(name='Arial',color='FF0EAA',bold=True,size=12); MUT=Font(name='Arial',color='8C93A6',size=9,italic=True); MINT=Font(name='Arial',color='9AD9C2',bold=True,size=10)
M='$#,##0'; P='0.0%'
def sheet(ws,title,widths):
    ws.sheet_view.showGridLines=False
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for r in range(1,60):
        for c in range(1,len(widths)+1): ws.cell(r,c).fill=DARK; ws.cell(r,c).font=INK
    ws['A1']=title; ws['A1'].font=Font(name='Arial',color='FFFFFF',bold=True,size=16)
    ws['A2']='Brighter Nights Lighting Technologies  ·  Lighting Partner Program  ·  rebuilt 2026-09-26 (100% fee credit, Site Value Evaluator ledger)'; ws['A2'].font=MUT
def hdr(ws,r,vals):
    for i,v in enumerate(vals,1): c=ws.cell(r,i,v); c.font=WH; c.fill=HEAD; c.alignment=Alignment(horizontal='center' if i>1 else 'left',vertical='center',wrap_text=True)
    ws.row_dimensions[r].height=30
# ---------- Assumptions
A=wb.active; A.title='Assumptions'; sheet(A,'Assumptions',[46,18,14,60])
hdr(A,4,['Fee tiers (marginal, like tax brackets)','Band width','Rate','Notes'])
tiers=[('First $250,000',250000,0.07),('$250,000 to $500,000',250000,0.06),('$500,000 to $750,000',250000,0.05),('$750,000 to $1,000,000',250000,0.04),('Above $1,000,000',10**9,0.03)]
for i,(n,w,r) in enumerate(tiers):
    A.cell(5+i,1,n); A.cell(5+i,2,w).number_format=M; c=A.cell(5+i,3,r); c.number_format=P; c.fill=INP
A['D5']='Rates are inputs. Change C5:C9 and every tab recalculates.'; A['D5'].font=MUT
hdr(A,11,['Program inputs','Value','','Notes'])
inputs=[('Reinvestment share of fee into partner ads',0.30,P,'Partner matches this amount'),('Fee credit rate toward ledger',1.00,P,'RULED 2026-09-26: 100% of every fee dollar'),('Trailing months of fee added at buyout',3,'0',''),('Soft cap, flag add ons projecting buyout past year',7,'0',''),('Add ons per year (light)',5000,M,'Routine maintenance never counts'),('Add ons per year (CRM level)',20000,M,'')]
for i,(n,v,f,note) in enumerate(inputs):
    A.cell(12+i,1,n); c=A.cell(12+i,2,v); c.number_format=f; c.fill=INP; A.cell(12+i,4,note).font=MUT
hdr(A,19,['Ledger (Site Value Evaluator)','Value','','Notes'])
ledger=[('Part 1  Market replacement cost of what was built',102650,'Base package, 13 items, midpoint of 2026 agency and freelance rates'),('Part 2  Contractor hours not spent teaching an agency',180,'hours'),('Contractor hourly value',150,'per hour'),('Part 3  Months of agency delay avoided',5,'months'),('Monthly gross used to price delay',41000,'about $492K a year'),('Share of gross driven by the site',0.35,'')]
for i,(n,v,note) in enumerate(ledger):
    A.cell(20+i,1,n); c=A.cell(20+i,2,v); c.fill=INP; c.number_format=P if isinstance(v,float) else (M if v>1000 else '0'); A.cell(20+i,4,note).font=MUT
A['A27']='Contractor time value'; A['B27']='=B21*B22'; A['B27'].number_format=M
A['A28']='Revenue not lost to delay'; A['B28']='=B23*B24*B25'; A['B28'].number_format=M
A['A29']='LEDGER OPENING VALUE'; A['A29'].font=PINK; A['B29']='=B20+B27+B28'; A['B29'].number_format=M; A['B29'].font=PINK
# ---------- Projection
Pj=wb.create_sheet('Projection'); sheet(Pj,'Five year projection and earn out',[34,16,16,16,16,16,16])
hdr(Pj,4,['','Year 1','Year 2','Year 3','Year 4','Year 5','Total'])
gross=[500000,750000,1000000,1300000,1600000]
Pj['A5']='Partner gross lighting revenue'
for i,g in enumerate(gross): c=Pj.cell(5,2+i,g); c.number_format=M; c.fill=INP
Pj['G5']='=SUM(B5:F5)'; Pj['G5'].number_format=M
Pj['A6']='Program fee (marginal tiers)'
for i in range(5):
    col=get_column_letter(2+i); g=f'{col}5'
    f=f"=MIN({g},Assumptions!$B$5)*Assumptions!$C$5+MIN(MAX({g}-Assumptions!$B$5,0),Assumptions!$B$6)*Assumptions!$C$6+MIN(MAX({g}-Assumptions!$B$5-Assumptions!$B$6,0),Assumptions!$B$7)*Assumptions!$C$7+MIN(MAX({g}-Assumptions!$B$5-Assumptions!$B$6-Assumptions!$B$7,0),Assumptions!$B$8)*Assumptions!$C$8+MAX({g}-Assumptions!$B$5-Assumptions!$B$6-Assumptions!$B$7-Assumptions!$B$8,0)*Assumptions!$C$9"
    Pj.cell(6,2+i,f).number_format=M
Pj['G6']='=SUM(B6:F6)'; Pj['G6'].number_format=M
rows=[('Blended fee rate','={c}6/{c}5',P,None),('Partner profit after fee','={c}5-{c}6',M,'=SUM(B{r}:F{r})'),('Our reinvestment into partner ads','={c}6*Assumptions!$B$12',M,'=SUM(B{r}:F{r})'),('Partner marketing match','={c}6*Assumptions!$B$12',M,'=SUM(B{r}:F{r})'),('Fee credit to ledger','={c}6*Assumptions!$B$13',M,'=SUM(B{r}:F{r})'),('Cumulative credit','=SUM($B$11:{c}11)',M,None),('Ledger balance (no add ons)','=Assumptions!$B$29',M,None),('Remaining to buy out','=MAX({c}13-{c}12,0)',M,None),('Buyout price today (remaining + trailing fee)','={c}14+{c}6*Assumptions!$B$14/12',M,None),('Owned outright?','=IF({c}12>={c}13,"YES","not yet")',None,None)]
for k,(n,f,fmt,tot) in enumerate(rows):
    r=7+k; Pj.cell(r,1,n)
    for i in range(5):
        col=get_column_letter(2+i); c=Pj.cell(r,2+i,f.format(c=col));
        if fmt: c.number_format=fmt
    if tot: Pj.cell(r,7,tot.format(r=r)).number_format=fmt
Pj['A16'].font=MINT
for i in range(5): Pj.cell(16,2+i).font=MINT
Pj['A18']='Earn out year (interpolated, no add ons)'; Pj['A18'].font=PINK
Pj['B18']='=IF(F12<Assumptions!B29,"past year 5",IF(B12>=Assumptions!B29,B12/B11,IF(C12>=Assumptions!B29,1+(Assumptions!B29-B12)/C11,IF(D12>=Assumptions!B29,2+(Assumptions!B29-C12)/D11,IF(E12>=Assumptions!B29,3+(Assumptions!B29-D12)/E11,4+(Assumptions!B29-E12)/F11)))))'; Pj['B18'].number_format='0.0'; Pj['B18'].font=PINK
Pj['A19']='Flat $492K gross, years to earn out'; Pj['B19']='=Assumptions!B29/(MIN(492000,Assumptions!B5)*Assumptions!C5+MIN(MAX(492000-Assumptions!B5,0),Assumptions!B6)*Assumptions!C6)/Assumptions!B13'; Pj['B19'].number_format='0.0'
Pj['A21']='With add ons (light pace)'; Pj['A21'].font=WH
Pj['A22']='Ledger by year end';
for i in range(5): Pj.cell(22,2+i,f'=Assumptions!$B$29+Assumptions!$B$16*{i+1}').number_format=M
Pj['A23']='Owned outright?'
for i in range(5): col=get_column_letter(2+i); Pj.cell(23,2+i,f'=IF({col}12>={col}22,"YES","not yet")')
Pj['A25']='With add ons (CRM level pace)'; Pj['A25'].font=WH
Pj['A26']='Ledger by year end'
for i in range(5): Pj.cell(26,2+i,f'=Assumptions!$B$29+Assumptions!$B$17*{i+1}').number_format=M
Pj['A27']='Owned outright?'
for i in range(5): col=get_column_letter(2+i); Pj.cell(27,2+i,f'=IF({col}12>={col}26,"YES","not yet")')
# ---------- Fee by gross
Fb=wb.create_sheet('Fee by gross'); sheet(Fb,'Fee at each gross level',[22,16,16,18])
hdr(Fb,4,['Annual gross','Fee','Blended rate','Profit after fee'])
for i,g in enumerate([250000,500000,750000,1000000,1250000,1500000,2000000]):
    r=5+i; Fb.cell(r,1,g).number_format=M
    f=f"=MIN(A{r},Assumptions!$B$5)*Assumptions!$C$5+MIN(MAX(A{r}-Assumptions!$B$5,0),Assumptions!$B$6)*Assumptions!$C$6+MIN(MAX(A{r}-Assumptions!$B$5-Assumptions!$B$6,0),Assumptions!$B$7)*Assumptions!$C$7+MIN(MAX(A{r}-Assumptions!$B$5-Assumptions!$B$6-Assumptions!$B$7,0),Assumptions!$B$8)*Assumptions!$C$8+MAX(A{r}-Assumptions!$B$5-Assumptions!$B$6-Assumptions!$B$7-Assumptions!$B$8,0)*Assumptions!$C$9"
    Fb.cell(r,2,f).number_format=M; Fb.cell(r,3,f'=B{r}/A{r}').number_format=P; Fb.cell(r,4,f'=A{r}-B{r}').number_format=M
# ---------- Agency comparison
Ag=wb.create_sheet('Agency comparison'); sheet(Ag,'Three year cost, program vs agency vs DIY',[40,18,18,18])
hdr(Ag,4,['','Lighting Partner Program','Typical agency','Do it yourself'])
Ag['A5']='Agency monthly retainer'; Ag['C5']=4000; Ag['C5'].number_format=M; Ag['C5'].fill=INP
Ag['A6']='Agency setup and site build'; Ag['C6']=12000; Ag['C6'].number_format=M; Ag['C6'].fill=INP
Ag['A7']='DIY build cost (Part 1 replacement)'; Ag['D7']='=Assumptions!B20'; Ag['D7'].number_format=M
Ag['A8']='Owner time teaching or building (Part 2)'; Ag['C8']='=Assumptions!B27'; Ag['D8']='=Assumptions!B27*2'
Ag['A9']='Revenue lost to delay (Part 3)'; Ag['C9']='=Assumptions!B28'; Ag['D9']='=Assumptions!B28*1.5'
for r in (8,9):
    for c in ('C','D'): Ag[f'{c}{r}'].number_format=M
Ag['A11']='Three year fees paid'; Ag['B11']='=SUM(Projection!B6:D6)'; Ag['C11']='=C5*36+C6'; Ag['D11']=0
Ag['A12']='Ads funded by provider over three years'; Ag['B12']='=SUM(Projection!B9:D9)'; Ag['C12']=0; Ag['D12']=0
Ag['A13']='Hidden costs (time and delay)'; Ag['B13']=0; Ag['C13']='=C8+C9'; Ag['D13']='=D7+D8+D9'
Ag['A14']='Net three year cost'; Ag['A14'].font=WH
for c in 'BCD': Ag[f'{c}14']=f'={c}11-{c}12+{c}13'; Ag[f'{c}14'].font=PINK
Ag['A15']='Credit toward owning the site'; Ag['B15']='=SUM(Projection!B11:D11)'; Ag['C15']=0; Ag['D15']=0
Ag['A16']='Asset owned at end (ledger credited)'; Ag['B16']='=MIN(B15,Assumptions!B29)'; Ag['C16']=0; Ag['D16']='=D7'
for r in range(11,17):
    for c in 'BCD': Ag[f'{c}{r}'].number_format=M
Ag['A18']='Program is cheaper than the agency by'; Ag['B18']='=C14-B14'; Ag['B18'].number_format=M; Ag['B18'].font=MINT
wb.save('Lighting_Partner_Program_Model.xlsx'); print('saved')
